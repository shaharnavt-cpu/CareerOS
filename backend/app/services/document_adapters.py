"""Safe byte-level adapters for Universal Listening documents."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Protocol
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError, PdfReadError

MAX_DOCUMENT_BYTES = 15 * 1024 * 1024
MAX_UPLOAD_FILES = 10


class DocumentAdapterError(ValueError):
    """An extraction failure with a safe, actionable user message."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.user_message = message


@dataclass(frozen=True)
class RawDocument:
    """Untrusted file bytes accepted by the listening queue."""

    filename: str
    content: bytes
    media_type: str | None = None

    @property
    def extension(self) -> str:
        return Path(self.filename).suffix.casefold()


@dataclass(frozen=True)
class ExtractedDocument:
    """Text plus format metadata produced by a document adapter."""

    text: str
    format: str
    adapter_version: str
    metadata: dict[str, Any] = field(default_factory=dict)


class DocumentAdapter(Protocol):
    """Extract inert text from one supported document format."""

    extensions: frozenset[str]
    version: str

    def extract(self, document: RawDocument) -> ExtractedDocument:
        """Extract text without executing embedded document content."""


def _require_content(text: str) -> str:
    normalized = text.replace("\x00", "").strip()
    if not normalized:
        raise DocumentAdapterError(
            "EMPTY_DOCUMENT",
            "This document does not contain readable text. Try exporting it again "
            "or add the details in the conversation.",
        )
    return normalized


class PdfDocumentAdapter:
    """Extract page-addressable text from PDFs using pypdf."""

    extensions = frozenset({".pdf"})
    version = "pdf-adapter-v1"

    def extract(self, document: RawDocument) -> ExtractedDocument:
        try:
            reader = PdfReader(io.BytesIO(document.content), strict=False)
            if reader.is_encrypted and reader.decrypt("") == 0:
                raise DocumentAdapterError(
                    "PASSWORD_PROTECTED_PDF",
                    "This PDF is password-protected. Remove the password and attach "
                    "it again.",
                )
            pages: list[str] = []
            for page_number, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    pages.append(f"[Page {page_number}]\n{text}")
        except DocumentAdapterError:
            raise
        except FileNotDecryptedError as exc:
            raise DocumentAdapterError(
                "PASSWORD_PROTECTED_PDF",
                "This PDF is password-protected. Remove the password and attach it "
                "again.",
            ) from exc
        except (PdfReadError, OSError, ValueError, TypeError) as exc:
            raise DocumentAdapterError(
                "CORRUPTED_PDF",
                "CareerOS could not read this PDF. Try opening and re-exporting it, "
                "then attach the new copy.",
            ) from exc

        return ExtractedDocument(
            text=_require_content("\n\n".join(pages)),
            format="PDF",
            adapter_version=self.version,
            metadata={"page_count": len(reader.pages)},
        )


class DocxDocumentAdapter:
    """Extract paragraphs and tables from modern Word documents."""

    extensions = frozenset({".docx"})
    version = "docx-adapter-v1"

    def extract(self, document: RawDocument) -> ExtractedDocument:
        try:
            word_document = Document(io.BytesIO(document.content))
            blocks = [
                paragraph.text.strip()
                for paragraph in word_document.paragraphs
                if paragraph.text.strip()
            ]
            table_rows = 0
            for table in word_document.tables:
                for row in table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        blocks.append(" | ".join(values))
                        table_rows += 1
        except (BadZipFile, PackageNotFoundError, KeyError, ValueError) as exc:
            raise DocumentAdapterError(
                "CORRUPTED_DOCX",
                "CareerOS could not read this Word document. Open it in Word or "
                "Google Docs, save a new .docx copy, and try again.",
            ) from exc

        return ExtractedDocument(
            text=_require_content("\n".join(blocks)),
            format="DOCX",
            adapter_version=self.version,
            metadata={
                "paragraph_count": len(word_document.paragraphs),
                "table_count": len(word_document.tables),
                "table_row_count": table_rows,
            },
        )


class TextDocumentAdapter:
    """Decode TXT, Markdown, and CSV without executing or interpreting markup."""

    extensions = frozenset({".txt", ".md", ".markdown", ".csv"})
    version = "text-document-adapter-v1"

    def extract(self, document: RawDocument) -> ExtractedDocument:
        text: str | None = None
        for encoding in ("utf-8-sig", "utf-8", "cp1252"):
            try:
                text = document.content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise DocumentAdapterError(
                "TEXT_ENCODING",
                "CareerOS could not read this text encoding. Save the file as UTF-8 "
                "and attach it again.",
            )

        metadata: dict[str, Any] = {}
        format_name = document.extension.removeprefix(".").upper()
        if document.extension == ".csv":
            try:
                rows = list(csv.reader(io.StringIO(text)))
            except csv.Error as exc:
                raise DocumentAdapterError(
                    "CORRUPTED_CSV",
                    "CareerOS could not read this CSV. Export a fresh CSV and try "
                    "again.",
                ) from exc
            text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
            metadata["row_count"] = len(rows)

        return ExtractedDocument(
            text=_require_content(text),
            format=format_name,
            adapter_version=self.version,
            metadata=metadata,
        )


class DocumentAdapterRegistry:
    """Extension point for local files and future external source adapters."""

    def __init__(self, adapters: list[DocumentAdapter] | None = None) -> None:
        configured = adapters or [
            PdfDocumentAdapter(),
            DocxDocumentAdapter(),
            TextDocumentAdapter(),
        ]
        self._adapters: dict[str, DocumentAdapter] = {}
        for adapter in configured:
            self.register(adapter)

    @property
    def supported_extensions(self) -> frozenset[str]:
        return frozenset(self._adapters)

    def register(self, adapter: DocumentAdapter) -> None:
        for extension in adapter.extensions:
            self._adapters[extension.casefold()] = adapter

    def validate(self, document: RawDocument) -> None:
        if len(document.content) > MAX_DOCUMENT_BYTES:
            raise DocumentAdapterError(
                "FILE_TOO_LARGE",
                "This file is larger than 15 MB. Attach a smaller export or split it "
                "into a few files.",
            )
        if document.extension not in self._adapters:
            raise DocumentAdapterError(
                "UNSUPPORTED_FILE",
                "CareerOS currently listens to PDF, DOCX, TXT, Markdown, and CSV "
                "files.",
            )

    def extract(self, document: RawDocument) -> ExtractedDocument:
        self.validate(document)
        return self._adapters[document.extension].extract(document)
