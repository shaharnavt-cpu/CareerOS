"""Universal Listening adapters, detection, queue, and progress API tests."""

from __future__ import annotations

import threading
import uuid
from contextlib import nullcontext
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy.exc import OperationalError

from app.api.listening import get_listening_queue
from app.main import app
from app.services.document_adapters import (
    DocumentAdapterError,
    DocumentAdapterRegistry,
    DocxDocumentAdapter,
    PdfDocumentAdapter,
    RawDocument,
)
from app.services.document_detection import DocumentDetector, DocumentType
from app.services.listening_queue import (
    JobStatus,
    ListeningQueue,
    PersistenceStatus,
    PipelineStage,
    ProcessingResult,
    StageStatus,
    UniversalListeningProcessor,
)


def make_pdf(text: str, *, password: str | None = None) -> bytes:
    """Create a small text PDF without adding a test-only rendering dependency."""

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    safe_text = text.replace("(", "[").replace(")", "]")
    stream.set_data(f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Professional Experience", level=1)
    document.add_paragraph("Software Engineer at Acme")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, SQL"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_pdf_ingestion_extracts_page_text_and_metadata() -> None:
    result = PdfDocumentAdapter().extract(
        RawDocument("resume.pdf", make_pdf("Resume Skills Python"), "application/pdf")
    )

    assert "Resume Skills Python" in result.text
    assert result.metadata["page_count"] == 1
    assert result.adapter_version == "pdf-adapter-v1"


def test_pdf_without_extractable_text_is_classified() -> None:
    with pytest.raises(DocumentAdapterError) as error:
        PdfDocumentAdapter().extract(RawDocument("scan.pdf", make_pdf("")))

    assert error.value.code == "NO_TEXT_FOUND"


def test_docx_ingestion_extracts_paragraphs_and_tables() -> None:
    result = DocxDocumentAdapter().extract(
        RawDocument(
            "career.docx",
            make_docx(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    )

    assert "Software Engineer at Acme" in result.text
    assert "Skills | Python, SQL" in result.text
    assert result.metadata["table_count"] == 1


def test_password_protected_pdf_has_actionable_error() -> None:
    with pytest.raises(DocumentAdapterError) as error:
        PdfDocumentAdapter().extract(
            RawDocument("locked.pdf", make_pdf("Private", password="secret"))
        )

    assert error.value.code == "PASSWORD_PROTECTED_PDF"
    assert "Remove the password" in error.value.user_message


@pytest.mark.parametrize(
    ("filename", "text", "expected"),
    [
        ("resume.pdf", "Skills Education Work Experience", DocumentType.RESUME),
        (
            "record.pdf",
            "Official Transcript Cumulative GPA Credit Hours",
            DocumentType.TRANSCRIPT,
        ),
        (
            "role.txt",
            "Job Description Responsibilities Qualifications Requirements",
            DocumentType.JOB_DESCRIPTION,
        ),
        (
            "award.docx",
            "Certificate awarded to Alex, credential ID 123",
            DocumentType.CERTIFICATE,
        ),
        (
            "paper.pdf",
            "Abstract Introduction Methodology References",
            DocumentType.RESEARCH_PAPER,
        ),
        ("work.pdf", "Portfolio Selected Work Case Study", DocumentType.PORTFOLIO),
        ("notes.txt", "Met Pat at the conference", DocumentType.GENERIC_DOCUMENT),
    ],
)
def test_document_detection(filename: str, text: str, expected: DocumentType) -> None:
    detection = DocumentDetector().detect(filename, text)

    assert detection.document_type is expected
    assert detection.reasons
    assert 0 <= detection.confidence <= 1


def test_queue_is_fifo_and_reports_progress() -> None:
    processed: list[str] = []
    first_started = threading.Event()
    release_first = threading.Event()

    def processor(job, report):
        report(PipelineStage.PARSE_DOCUMENT, "Reading")
        processed.append(job.document.filename)
        if job.document.filename == "first.txt":
            first_started.set()
            release_first.wait(timeout=2)
        report(PipelineStage.GENERATE_PROPOSALS, "Preparing")
        return ProcessingResult(
            source_id=uuid.uuid4(),
            source_type="RESUME",
            document_type="RESUME",
            detection_confidence=0.9,
            evidence_count=2,
        )

    listening_queue = ListeningQueue(processor)
    first = listening_queue.enqueue(RawDocument("first.txt", b"first"))
    second = listening_queue.enqueue(RawDocument("second.txt", b"second"))
    assert first_started.wait(timeout=2)
    assert listening_queue.get(second.id).status is JobStatus.QUEUED
    release_first.set()
    assert listening_queue.join()

    assert processed == ["first.txt", "second.txt"]
    assert listening_queue.get(first.id).status is JobStatus.COMPLETE
    assert (
        listening_queue.get(second.id)
        .stage_timings[PipelineStage.GENERATE_PROPOSALS]
        .status
        is StageStatus.COMPLETE
    )
    listening_queue.shutdown()


def test_failed_queue_job_can_be_retried() -> None:
    attempts = 0

    def processor(job, report):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise DocumentAdapterError(
                "CORRUPTED_PDF", "Export a fresh PDF and try again."
            )
        report(PipelineStage.PARSE_DOCUMENT, "Reading")
        return ProcessingResult(
            source_id=uuid.uuid4(),
            source_type="RESUME",
            document_type="RESUME",
            detection_confidence=0.9,
            evidence_count=1,
        )

    listening_queue = ListeningQueue(processor)
    job = listening_queue.enqueue(RawDocument("resume.pdf", b"pdf"))
    assert listening_queue.join()
    assert listening_queue.get(job.id).status is JobStatus.FAILED
    assert listening_queue.get(job.id).failure.code == "PARSER_FAILED"

    retried = listening_queue.retry(job.id)
    assert retried.attempt == 2
    assert listening_queue.join()
    assert listening_queue.get(job.id).status is JobStatus.COMPLETE
    listening_queue.shutdown()


def test_database_timeout_retains_parsing_and_resumes_at_persistence() -> None:
    class CountingAdapters(DocumentAdapterRegistry):
        def __init__(self) -> None:
            super().__init__()
            self.extractions = 0

        def extract(self, document):
            self.extractions += 1
            return super().extract(document)

    class RecoveringRepository:
        def __init__(self) -> None:
            self.persistence_attempts = 0
            self.evidence = []

        def persist_ingestion(self, session, source, evidence, clusters):
            del session, clusters
            self.persistence_attempts += 1
            if self.persistence_attempts == 1:
                raise OperationalError(
                    "connect",
                    {},
                    TimeoutError("connection timed out to localhost:5432"),
                )
            self.evidence = evidence
            return SimpleNamespace(
                id=source.id,
                source_type=source.source_type,
                evidence_count=len(evidence),
            )

        def list_evidence(self, session):
            del session
            return self.evidence

        def persist_clusters(self, session, user_id, clusters):
            del session, user_id, clusters

        def persist_proposal_links(self, session, user_id, proposals, clusters):
            del session, user_id, proposals, clusters

    adapters = CountingAdapters()
    repository = RecoveringRepository()
    processor = UniversalListeningProcessor(
        lambda: nullcontext(object()),
        adapters=adapters,
        repository=repository,
    )
    listening_queue = ListeningQueue(processor)
    job = listening_queue.enqueue(
        RawDocument("resume.txt", b"Experience\nSoftware Engineer at Acme")
    )
    assert listening_queue.join()

    failed = listening_queue.get(job.id)
    assert failed.status is JobStatus.FAILED
    assert failed.current_stage is PipelineStage.PERSIST_EVIDENCE
    assert failed.failure.code == "DATABASE_TIMEOUT"
    assert failed.extracted_character_count > 0
    assert failed.evidence_count > 0
    assert failed.persistence_status is PersistenceStatus.FAILED
    assert adapters.extractions == 1

    listening_queue.retry(job.id)
    assert listening_queue.join()
    recovered = listening_queue.get(job.id)
    assert recovered.status is JobStatus.COMPLETE
    assert recovered.persistence_status is PersistenceStatus.SAVED
    assert recovered.retry_count == 1
    assert adapters.extractions == 1
    assert (
        sum(log["event"] == "PARSE_DOCUMENT" for log in recovered.developer_logs) == 1
    )
    listening_queue.shutdown()


def test_multiple_upload_and_progress_endpoints() -> None:
    def processor(job, report):
        report(PipelineStage.CREATE_EVIDENCE, "Extracting")
        return ProcessingResult(
            source_id=uuid.uuid4(),
            source_type="GENERIC_DOCUMENT",
            document_type="GENERIC_DOCUMENT",
            detection_confidence=0.5,
            evidence_count=1,
        )

    listening_queue = ListeningQueue(processor)
    original = dict(app.dependency_overrides)
    app.dependency_overrides[get_listening_queue] = lambda: listening_queue
    try:
        with TestClient(app) as client:
            response = client.post(
                "/api/listening/uploads",
                files=[
                    ("files", ("one.txt", b"one", "text/plain")),
                    ("files", ("two.md", b"two", "text/markdown")),
                ],
            )
            assert response.status_code == 202
            assert len(response.json()["jobs"]) == 2
            assert listening_queue.join()

            jobs = client.get("/api/listening/jobs")
            assert jobs.status_code == 200
            assert {job["status"] for job in jobs.json()} == {"COMPLETE"}
            job_id = jobs.json()[0]["id"]
            progress = client.get(f"/api/listening/jobs/{job_id}")
            assert progress.json()["current_stage"] == "COMPLETE"
            assert progress.json()["source_id"]
            assert progress.json()["diagnostics"] is None

            developer = client.get(
                f"/api/listening/jobs/{job_id}?developer=true"
            ).json()
            assert developer["diagnostics"]["developer_logs"]
            assert developer["stages"]
    finally:
        app.dependency_overrides = original
        listening_queue.shutdown()
